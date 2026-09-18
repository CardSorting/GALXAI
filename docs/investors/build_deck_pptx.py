import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette
    COLOR_BG = RGBColor(6, 8, 13)          # #06080D
    COLOR_CARD = RGBColor(13, 18, 30)       # #0D121E
    COLOR_CARD_BORDER = RGBColor(30, 41, 59)# #1E293B
    COLOR_CYAN = RGBColor(0, 240, 255)      # #00F0FF
    COLOR_PURPLE = RGBColor(168, 85, 247)   # #A855F7
    COLOR_PINK = RGBColor(236, 72, 153)     # #EC4899
    COLOR_EMERALD = RGBColor(16, 185, 129)  # #10B981
    COLOR_AMBER = RGBColor(245, 158, 11)    # #F59E0B
    COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252) # #F8FAFC
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8

    def set_slide_bg(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, slide_num, total_slides=10):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.5))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        
        r1 = p.add_run()
        r1.text = "GALX "
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = COLOR_CYAN
        
        r2 = p.add_run()
        r2.text = "•  Venture Formation Thesis"
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_TEXT_MUTED

        num_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.4), Inches(2.0), Inches(0.5))
        tf2 = num_box.text_frame
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.RIGHT
        r3 = p2.add_run()
        r3.text = f"Slide {slide_num:02d} / {total_slides:02d}"
        r3.font.size = Pt(11)
        r3.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(slide, left, top, width, height, title, body, border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = tf.margin_bottom = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_TEXT_PRIMARY

        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(8)
        return card

    slides_data = [
        ("Sustainable, Abundant AI Access for Anyone Who Wants to Build",
         "The abstraction layer for the next era of software creation. Eliminating infrastructure management so developers can focus entirely on what they want to build.",
         [("Two Managed Models", "Concentrated surface, including GALX Nova"),
          ("Standard API Access", "Instant drop-in integration"),
          ("Abundant Serving", "Reliable operational execution"),
          ("Build What You Want", "Simplicity is the product decision")]),

        ("The Future: Abundant Intelligence, Fragmented Infrastructure",
         "AI capability is becoming abundant, while the systems required to access and sustain it are fragmenting.",
         [("What Is Coming", "AI capability is becoming the fundamental substrate of all software creation."),
          ("The Impending Bottleneck", "Multiplying models, providers, and pricing structures create severe operational fragmentation.")]),

        ("The Non-Consensus Thesis: The Sufficiency Principle",
         "For a large class of applications, developers do not need maximum available intelligence. They need sufficient intelligence, reliably and economically delivered, to complete the task.",
         [("The Consensus Trap", "Obsession with benchmark supremacy and endless model dropdowns."),
          ("The Intermediary Moat", "Model proliferation creates an opportunity for an intermediary to absorb complexity.")]),

        ("Product Today: Radical Simplicity",
         "Two models. API access. Abundant serving. Build what you want.",
         [("Two Models", "Access focused models under clear GALX identities such as GALX Nova."),
          ("Standard API", "Drop-in execution with zero custom proxy plumbing."),
          ("Abundant Serving", "GALX handles availability, capacity, and resilience behind the scenes.")]),

        ("The Wedge: Builders Who Want to Make Something",
         "We focus on high-velocity creators who want to build software rather than manage inference infrastructure.",
         [("Independent Developers", "Building indie SaaS and productivity tools with predictable costs."),
          ("Creative Technologists", "Crafting interactive apps and media without cloud console friction."),
          ("AI Experimenters", "Testing application loops and agent workflows rapidly.")]),

        ("Why This Can Become Enormous",
         "GALX sits directly between the application and fragmented intelligence infrastructure.",
         [("The Expansion Loop", "More Builders -> More Applications -> More Application Usage -> More Execution Through GALX."),
          ("The Persistent Execution Layer", "Software built on GALX generates ongoing recurring execution demand with every user click.")]),

        ("Asymmetric Advantage: Absorb Complexity, Sell Simplicity",
         "Operating knowledge turns heterogeneous, volatile AI capacity into a reliable developer product.",
         [("Developer Simplicity", "Clean, dependable API with zero infrastructure maintenance."),
          ("Proprietary Depth", "Proprietary operational knowledge required to run heterogeneous capacity economically.")]),

        ("Founder-Market Fit: Built From Direct Systems Operation",
         "Arrived at the problem through building AI applications and living the operational constraints.",
         [("Hands-On Experience", "Direct exposure to building AI applications and operating autonomous agent systems."),
          ("Operational Learning", "Learning and iterating through live traffic and operational feedback.")]),

        ("Evidence: Observations vs. Next Validations",
         "Separating conviction from verified observations and upcoming milestones.",
         [("What We Believe", "Abundant intelligence belongs behind a dependable abstraction layer."),
          ("What We Have Observed", "Representative workloads require substantially less model capability than assumed."),
          ("What We Demonstrated", "A functional, high-reliability platform serving real developer requests."),
          ("What We Are Testing Next", "Builder retention, application usage growth, and long-term execution margins.")]),

        ("The Venture Test: What Becomes Enormous?",
         "GALX becomes the primary abstraction and execution layer through which software consumes AI capability.",
         [("Scale Capacity", "Expand serving infrastructure to support high-concurrency builder apps."),
          ("Instrument Resilience", "Deepen operational knowledge tooling and serving cost optimization."),
          ("Onboard Builders", "Seed the next 1,000 applications building on top of GALX execution.")])
    ]

    for idx, (title, subtitle, cards) in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank_slide_layout)
        set_slide_bg(slide)
        add_header(slide, idx, len(slides_data))

        # Title & Subtitle box
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.733), Inches(1.4))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(24)
        p.font.color.rgb = COLOR_TEXT_PRIMARY

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = COLOR_CYAN
        p_sub.space_before = Pt(6)

        # Cards Layout
        num_cards = len(cards)
        card_w = (11.733 - (0.3 * (num_cards - 1))) / num_cards
        for c_idx, (c_title, c_body) in enumerate(cards):
            c_left = 0.8 + c_idx * (card_w + 0.3)
            add_card(slide, Inches(c_left), Inches(2.8), Inches(card_w), Inches(3.6), c_title, c_body)

    out_path = "/Users/bozoegg/Desktop/GALXAI/docs/investors/GALXAI_Pitch_Deck.pptx"
    prs.save(out_path)
    print(f"Saved pitch deck to {out_path}")

if __name__ == "__main__":
    create_deck()
